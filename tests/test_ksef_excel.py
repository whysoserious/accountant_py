#!/usr/bin/env python3
"""
Tests for the accountant Excel report.

Everything here runs offline. Workbook assembly is exercised with no API access
at all; the description step is driven through a stub client so the prompt
contract can be asserted without spending a token.
"""

import json
import logging
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import ksef_excel
from ksef_excel import (
    COLUMN_HEADERS,
    MANIFEST_FILENAME,
    build_workbook,
    deduplicate,
    describe_invoices,
    filter_by_month,
    format_date,
    format_positions,
    load_records,
    mark_non_deductible,
    match_non_deductible,
    manifest_key,
    parse_invoice_xml,
    write_report,
)

from tests.fixtures.synthetic import (
    BUYER_NAME,
    BUYER_NIP,
    SELLER_NAME,
    SELLER_NIP,
    build_fa3_xml,
    build_ksef_pdf,
    build_text_pdf,
    default_positions,
    mixed_rate_positions,
    synthetic_ksef_number,
)

KSEF_NUMBER = "1000000000-20260625-AAAAAAAAAAAA-01"


class StubMessage:
    def __init__(self, text):
        self.content = [type("Block", (), {"text": text})()]


class StubMessages:
    """Records every call so the prompt contract can be asserted."""

    def __init__(self, replies=None, error=None):
        self.replies = list(replies or [])
        self.error = error
        self.calls = []

    def create(self, **kwargs):
        self.calls.append(kwargs)
        if self.error is not None:
            raise self.error
        return StubMessage(self.replies.pop(0) if self.replies else "opis")


class StubClient:
    def __init__(self, replies=None, error=None):
        self.messages = StubMessages(replies=replies, error=error)


class TestColumnLayout:
    def test_headers_match_the_accountants_sheet_then_the_two_additions(self):
        assert COLUMN_HEADERS == [
            "NIP",
            "Kontrahent",
            "Numer KSeF",
            "Numer dokumentu",
            "Data wystawienia",
            "Netto",
            "VAT",
            "Brutto",
            "Waluta",
            "Opis faktury - czego dotyczy w działalnosci",
            "Nazwa pliku",
            "pozycje",
        ]

    def test_there_are_exactly_twelve_columns(self):
        assert len(COLUMN_HEADERS) == 12


class TestParseInvoiceXml:
    def test_extracts_the_scalar_fields(self):
        record = parse_invoice_xml(build_fa3_xml())

        assert record.document_number == "FS 1/2026"
        assert record.issue_date == "2026-06-25"
        assert record.currency == "PLN"

    def test_realistic_document_carries_no_ksef_number(self):
        """
        The regression this pins: real FA(3) documents have no KSeF number, so
        parsing alone leaves the column empty and load_records must recover it.
        """
        assert parse_invoice_xml(build_fa3_xml()).ksef_number == ""

    def test_reads_the_ksef_number_when_the_element_is_present(self):
        record = parse_invoice_xml(build_fa3_xml(ksef_number=KSEF_NUMBER, include_ksef_number=True))
        assert record.ksef_number == KSEF_NUMBER

    def test_captures_both_parties_regardless_of_role(self):
        record = parse_invoice_xml(build_fa3_xml(), role="buyer")
        assert record.seller_nip == SELLER_NIP
        assert record.buyer_nip == BUYER_NIP

    def test_is_namespace_agnostic(self):
        bare = parse_invoice_xml(build_fa3_xml(ksef_number=KSEF_NUMBER))
        prefixed = parse_invoice_xml(build_fa3_xml(ksef_number=KSEF_NUMBER, prefix="tns:"))

        assert prefixed.ksef_number == bare.ksef_number
        assert prefixed.counterparty == bare.counterparty
        assert prefixed.net == bare.net
        assert len(prefixed.positions) == len(bare.positions)

    def test_buyer_role_reports_the_seller_as_counterparty(self):
        record = parse_invoice_xml(build_fa3_xml(), role="buyer")
        assert record.counterparty == SELLER_NAME
        assert record.nip == SELLER_NIP

    def test_seller_role_reports_the_buyer_as_counterparty(self):
        record = parse_invoice_xml(build_fa3_xml(), role="seller")
        assert record.counterparty == BUYER_NAME
        assert record.nip == BUYER_NIP

    def test_sums_single_rate_totals(self):
        record = parse_invoice_xml(build_fa3_xml(positions=default_positions()))

        assert record.net == Decimal("200.00")
        assert record.vat == Decimal("46.00")
        assert record.gross == Decimal("246.00")

    def test_sums_across_every_vat_rate(self):
        """P_13_x / P_14_x pairs must all be added, not just the first."""
        record = parse_invoice_xml(build_fa3_xml(positions=mixed_rate_positions()))

        assert record.net == Decimal("300.00")
        assert record.vat == Decimal("33.00")
        assert record.gross == Decimal("333.00")

    def test_falls_back_to_gross_minus_net_without_rate_totals(self):
        record = parse_invoice_xml(
            build_fa3_xml(positions=default_positions(), include_rate_totals=False)
        )

        assert record.gross == Decimal("246.00")
        assert record.vat == Decimal("46.00")

    def test_parses_line_items(self):
        record = parse_invoice_xml(build_fa3_xml(positions=default_positions()))

        assert len(record.positions) == 2
        first, second = record.positions
        assert first.name == "Usluga hostingu"
        assert first.unit == "szt."
        assert first.quantity == Decimal("1")
        assert first.net == Decimal("100.00")
        assert first.vat_rate == "23"
        assert second.name == "Domena .pl"
        assert second.quantity == Decimal("2")

    def test_missing_ksef_number_yields_empty_string_not_a_crash(self):
        record = parse_invoice_xml(build_fa3_xml(include_ksef_number=False))
        assert record.ksef_number == ""

    def test_invoice_without_line_items_still_parses(self):
        record = parse_invoice_xml(build_fa3_xml(positions=[]))
        assert record.positions == []
        assert record.net == Decimal("0.00")

    def test_malformed_xml_raises_value_error(self):
        with pytest.raises(ValueError):
            parse_invoice_xml(b"<Faktura><unclosed>")

    def test_description_starts_empty(self):
        record = parse_invoice_xml(build_fa3_xml())
        assert record.description == ""


class TestFormatPositions:
    def test_numbers_each_item_on_its_own_line(self):
        record = parse_invoice_xml(build_fa3_xml(positions=default_positions()))
        text = format_positions(record.positions)

        lines = text.split("\n")
        assert len(lines) == 2
        assert (
            lines[0]
            == "1. Usluga hostingu — 1 szt. × 100.00 = netto 100.00, VAT 23%, brutto 123.00"
        )
        assert lines[1] == "2. Domena .pl — 2 szt. × 50.00 = netto 100.00, VAT 23%, brutto 123.00"

    def test_empty_positions_give_an_empty_cell(self):
        assert format_positions([]) == ""

    def test_renders_non_numeric_vat_rate_verbatim(self):
        """Exempt and out-of-scope invoices carry codes like 'zw' rather than a number."""
        from tests.fixtures.synthetic import Position

        xml = build_fa3_xml(positions=[Position(name="Usluga medyczna", vat_rate="zw")])
        record = parse_invoice_xml(xml)
        text = format_positions(record.positions)

        assert "VAT zw" in text
        assert "%" not in text

    def test_strips_trailing_zeros_from_quantity(self):
        from tests.fixtures.synthetic import Position

        xml = build_fa3_xml(positions=[Position(name="Usluga", quantity="1.000")])
        record = parse_invoice_xml(xml)
        assert "1 szt." in format_positions(record.positions)


class TestFormatDate:
    def test_converts_iso_to_polish_display_order(self):
        assert format_date("2026-06-25") == "25.06.2026"

    def test_tolerates_a_datetime_value(self):
        assert format_date("2026-06-25T10:30:00Z") == "25.06.2026"

    def test_passes_through_an_unparseable_value(self):
        assert format_date("nieznana") == "nieznana"

    def test_passes_through_empty(self):
        assert format_date("") == ""


class TestLoadRecords:
    def write_pair(self, directory, stem, xml_bytes, with_pdf=True):
        directory.mkdir(parents=True, exist_ok=True)
        (directory / f"{stem}.xml").write_bytes(xml_bytes)
        if with_pdf:
            (directory / f"{stem}.pdf").write_bytes(b"%PDF-1.4 stub")

    def test_reads_every_xml_under_the_directory(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_pair(out, "2026-06-25, A, FS 1, ksef", build_fa3_xml(invoice_number="FS 1"))
        self.write_pair(out, "2026-06-26, B, FS 2, ksef", build_fa3_xml(invoice_number="FS 2"))

        records = load_records(str(tmp_path), logger=logger)

        assert {r.document_number for r in records} == {"FS 1", "FS 2"}

    def test_derives_the_local_pdf_filename_from_the_sidecar(self, tmp_path, logger):
        out = tmp_path / "output"
        stem = "2026-06-25, Dostawca, FS 1, ksef"
        self.write_pair(out, stem, build_fa3_xml())

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].filename == f"{stem}.pdf"

    def test_keeps_the_row_and_names_the_xml_when_no_pdf_exists(self, tmp_path, logger):
        """
        PDF rendering fails for some invoices, so there is genuinely no PDF to
        name. Reporting the XML that does exist beats naming a file the
        accountant cannot open, and both beat dropping the invoice.
        """
        out = tmp_path / "output"
        stem = "2026-06-25, Dostawca, FS 1, ksef"
        self.write_pair(out, stem, build_fa3_xml(), with_pdf=False)

        records = load_records(str(tmp_path), logger=logger)

        assert len(records) == 1
        assert records[0].filename == f"{stem}.xml"

    def test_skips_an_unparseable_xml_and_continues(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_pair(out, "good", build_fa3_xml(invoice_number="FS OK"))
        (out / "broken.xml").write_bytes(b"<Faktura><unclosed>")

        records = load_records(str(tmp_path), logger=logger)

        assert len(records) == 1
        assert records[0].document_number == "FS OK"

    def test_returns_empty_list_for_a_directory_with_no_xml(self, tmp_path, logger):
        (tmp_path / "output").mkdir()
        assert load_records(str(tmp_path), logger=logger) == []

    def test_returns_empty_list_for_a_missing_directory(self, tmp_path, logger):
        assert load_records(str(tmp_path / "absent"), logger=logger) == []

    def test_orders_records_by_issue_date(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_pair(out, "b", build_fa3_xml(issue_date="2026-06-26", invoice_number="second"))
        self.write_pair(out, "a", build_fa3_xml(issue_date="2026-06-24", invoice_number="first"))

        records = load_records(str(tmp_path), logger=logger)

        assert [r.document_number for r in records] == ["first", "second"]


class TestFilterByMonth:
    def make(self, issue_date):
        return parse_invoice_xml(build_fa3_xml(issue_date=issue_date))

    def test_keeps_only_the_requested_month(self):
        records = [self.make("2026-06-25"), self.make("2026-07-01"), self.make("2026-05-31")]
        assert len(filter_by_month(records, "2026-06")) == 1

    def test_none_keeps_everything(self):
        records = [self.make("2026-06-25"), self.make("2026-07-01")]
        assert len(filter_by_month(records, None)) == 2


class TestBuildWorkbook:
    def test_first_row_is_the_header(self):
        records = [parse_invoice_xml(build_fa3_xml())]
        sheet = build_workbook(records).active

        assert [c.value for c in sheet[1]] == COLUMN_HEADERS

    def test_writes_one_row_per_record(self):
        records = [
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 1")),
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 2")),
        ]
        sheet = build_workbook(records).active

        assert sheet.max_row == 3

    def test_money_cells_are_numeric_so_the_column_can_be_summed(self):
        records = [parse_invoice_xml(build_fa3_xml(positions=default_positions()))]
        sheet = build_workbook(records).active

        netto, vat, brutto = sheet.cell(2, 6), sheet.cell(2, 7), sheet.cell(2, 8)
        for cell in (netto, vat, brutto):
            assert isinstance(cell.value, (int, float)), f"{cell.value!r} is not numeric"
        assert netto.value == pytest.approx(200.00)
        assert vat.value == pytest.approx(46.00)
        assert brutto.value == pytest.approx(246.00)

    def test_date_is_written_in_polish_display_order(self):
        records = [parse_invoice_xml(build_fa3_xml(issue_date="2026-06-25"))]
        sheet = build_workbook(records).active
        assert sheet.cell(2, 5).value == "25.06.2026"

    def test_description_column_is_empty_without_the_ai_step(self):
        records = [parse_invoice_xml(build_fa3_xml())]
        sheet = build_workbook(records).active
        assert sheet.cell(2, 10).value in (None, "")

    def test_pozycje_cell_wraps_and_holds_every_item(self):
        records = [parse_invoice_xml(build_fa3_xml(positions=default_positions()))]
        sheet = build_workbook(records).active

        cell = sheet.cell(2, 12)
        assert "1. Usluga hostingu" in cell.value
        assert "2. Domena .pl" in cell.value
        assert cell.alignment.wrap_text is True

    def test_filename_column_carries_the_local_name(self, tmp_path, logger):
        out = tmp_path / "output"
        out.mkdir()
        stem = "2026-06-25, Dostawca, FS 1, ksef"
        (out / f"{stem}.xml").write_bytes(build_fa3_xml())
        (out / f"{stem}.pdf").write_bytes(b"%PDF")

        records = load_records(str(tmp_path), logger=logger)
        sheet = build_workbook(records).active

        assert sheet.cell(2, 11).value == f"{stem}.pdf"

    def test_empty_record_list_still_produces_a_header(self):
        sheet = build_workbook([]).active
        assert [c.value for c in sheet[1]] == COLUMN_HEADERS
        assert sheet.max_row == 1


class TestWriteReport:
    def test_writes_a_workbook_that_reads_back(self, tmp_path):
        records = [parse_invoice_xml(build_fa3_xml(positions=default_positions()))]
        target = tmp_path / "nested" / "ksef-2026-06.xlsx"

        write_report(records, str(target))

        assert target.exists()
        sheet = load_workbook(str(target)).active
        assert [c.value for c in sheet[1]] == COLUMN_HEADERS
        assert sheet.cell(2, 6).value == pytest.approx(200.00)


class TestDescribeInvoices:
    def test_fills_the_description_column(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["usluga hostingu na potrzeby dzialalnosci"])

        describe_invoices(records, client, logger)

        assert records[0].description == "usluga hostingu na potrzeby dzialalnosci"

    def test_sends_the_counterparty_positions_and_totals(self, logger):
        records = [parse_invoice_xml(build_fa3_xml(positions=default_positions()))]
        client = StubClient(replies=["opis"])

        describe_invoices(records, client, logger)

        prompt = client.messages.calls[0]["messages"][0]["content"]
        assert SELLER_NAME in prompt
        assert "Usluga hostingu" in prompt
        assert "200.00" in prompt

    def test_does_not_send_the_raw_xml(self, logger):
        """Only the fields the task needs should leave the machine."""
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["opis"])

        describe_invoices(records, client, logger)

        prompt = client.messages.calls[0]["messages"][0]["content"]
        assert "<Faktura" not in prompt
        assert "NumerKSeFDokumentu" not in prompt

    def test_one_failure_leaves_that_cell_empty_and_continues(self, logger):
        records = [
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 1")),
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 2")),
        ]
        client = StubClient(error=RuntimeError("api unavailable"))

        describe_invoices(records, client, logger)

        assert all(r.description == "" for r in records)
        assert len(client.messages.calls) == 2, "a failure must not abort the remaining rows"

    def test_strips_whitespace_from_the_reply(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["  opis z bialymi znakami  \n"])

        describe_invoices(records, client, logger)

        assert records[0].description == "opis z bialymi znakami"

    def test_uses_the_project_sonnet_model(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["opis"])

        describe_invoices(records, client, logger)

        assert client.messages.calls[0]["model"] == ksef_excel.CLAUDE_SONNET_MODEL

    def test_no_records_makes_no_calls(self, logger):
        client = StubClient()
        describe_invoices([], client, logger)
        assert client.messages.calls == []


class TestKsefNumberResolution:
    """
    Real FA(3) documents carry no KSeF reference number, so it is recovered
    from the manifest written at download time, or failing that from the text
    of the rendered PDF.
    """

    def write_invoice(self, directory, stem, xml_bytes, pdf_bytes=None):
        directory.mkdir(parents=True, exist_ok=True)
        (directory / f"{stem}.xml").write_bytes(xml_bytes)
        if pdf_bytes is not None:
            (directory / f"{stem}.pdf").write_bytes(pdf_bytes)

    def write_manifest(self, directory, entries):
        directory.mkdir(parents=True, exist_ok=True)
        (directory / MANIFEST_FILENAME).write_text(json.dumps(entries), encoding="utf-8")

    def test_manifest_key_is_stable_across_renames(self):
        """Keyed on invoice identity, so moving or renaming files cannot orphan it."""
        first = manifest_key(SELLER_NIP, "FS 1/2026", "2026-06-25")
        second = manifest_key(f" {SELLER_NIP} ", " FS 1/2026 ", " 2026-06-25 ")
        assert first == second

    def test_resolves_from_the_manifest(self, tmp_path, logger):
        number = synthetic_ksef_number()
        out = tmp_path / "output"
        self.write_invoice(out, "renamed-by-a-later-step", build_fa3_xml())
        self.write_manifest(out, {manifest_key(SELLER_NIP, "FS 1/2026", "2026-06-25"): number})

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == number

    def test_manifest_survives_the_file_being_renamed(self, tmp_path, logger):
        """The manifest lives in the month directory; the invoice moved to output/."""
        number = synthetic_ksef_number()
        month = tmp_path / "2026-06"
        self.write_manifest(month, {manifest_key(SELLER_NIP, "FS 1/2026", "2026-06-25"): number})
        out = tmp_path / "output"
        self.write_invoice(out, "2026-06-25, Dostawca, FS 1-2026, ksef", build_fa3_xml())

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == number

    def test_falls_back_to_the_pdf_text(self, tmp_path, logger):
        number = synthetic_ksef_number()
        out = tmp_path / "output"
        self.write_invoice(out, "legacy", build_fa3_xml(), build_ksef_pdf(number))

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == number

    def test_manifest_wins_over_the_pdf(self, tmp_path, logger):
        """The manifest comes straight from the API; PDF text is a recovery path."""
        authoritative = synthetic_ksef_number(date="20260101")
        from_pdf = synthetic_ksef_number(date="20261231")
        out = tmp_path / "output"
        self.write_invoice(out, "both", build_fa3_xml(), build_ksef_pdf(from_pdf))
        self.write_manifest(
            out, {manifest_key(SELLER_NIP, "FS 1/2026", "2026-06-25"): authoritative}
        )

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == authoritative

    def test_embedded_element_wins_over_everything(self, tmp_path, logger):
        in_xml = synthetic_ksef_number(date="20260202")
        out = tmp_path / "output"
        self.write_invoice(
            out,
            "embedded",
            build_fa3_xml(ksef_number=in_xml, include_ksef_number=True),
            build_ksef_pdf(synthetic_ksef_number(date="20261231")),
        )

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == in_xml

    def test_stays_empty_when_no_source_has_it(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_invoice(out, "nothing", build_fa3_xml())

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == ""

    def test_a_pdf_without_a_number_does_not_crash(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_invoice(out, "plain", build_fa3_xml(), build_text_pdf("Faktura bez numeru"))

        records = load_records(str(tmp_path), logger=logger)

        assert records[0].ksef_number == ""

    def test_a_corrupt_pdf_does_not_crash(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_invoice(out, "broken", build_fa3_xml(), b"%PDF-1.4 truncated garbage")

        records = load_records(str(tmp_path), logger=logger)

        assert len(records) == 1
        assert records[0].ksef_number == ""

    def test_unreadable_manifest_is_survivable(self, tmp_path, logger):
        out = tmp_path / "output"
        self.write_invoice(out, "inv", build_fa3_xml())
        (out / MANIFEST_FILENAME).write_text("{not json", encoding="utf-8")

        records = load_records(str(tmp_path), logger=logger)

        assert len(records) == 1

    def test_manifests_from_several_months_are_merged(self, tmp_path, logger):
        june = synthetic_ksef_number(date="20260625")
        july = synthetic_ksef_number(date="20260710")
        self.write_manifest(
            tmp_path / "2026-06", {manifest_key(SELLER_NIP, "FS 6", "2026-06-25"): june}
        )
        self.write_manifest(
            tmp_path / "2026-07", {manifest_key(SELLER_NIP, "FS 7", "2026-07-10"): july}
        )
        out = tmp_path / "output"
        self.write_invoice(out, "a", build_fa3_xml(invoice_number="FS 6", issue_date="2026-06-25"))
        self.write_invoice(out, "b", build_fa3_xml(invoice_number="FS 7", issue_date="2026-07-10"))

        found = {
            r.document_number: r.ksef_number for r in load_records(str(tmp_path), logger=logger)
        }

        assert found == {"FS 6": june, "FS 7": july}

    def test_the_column_carries_the_resolved_number(self, tmp_path, logger):
        number = synthetic_ksef_number()
        out = tmp_path / "output"
        self.write_invoice(out, "inv", build_fa3_xml(), build_ksef_pdf(number))

        sheet = build_workbook(load_records(str(tmp_path), logger=logger)).active

        assert sheet.cell(2, 3).value == number


class TestDescribeProgressLogging:
    def test_logs_one_line_per_invoice_with_a_running_count(self, caplog):
        records = [
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 1")),
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 2")),
            parse_invoice_xml(build_fa3_xml(invoice_number="FS 3")),
        ]
        client = StubClient(replies=["paliwo", "paliwo", "paliwo"])
        log = logging.getLogger("progress-test")

        with caplog.at_level(logging.INFO, logger="progress-test"):
            describe_invoices(records, client, log, categories=["paliwo"])

        text = caplog.text
        for i in (1, 2, 3):
            assert f"[{i}/3]" in text, f"missing progress line for invoice {i}"
        assert "Classification complete: 3 written" in text

    def test_a_failure_is_logged_against_its_own_index(self, caplog):
        records = [parse_invoice_xml(build_fa3_xml(invoice_number="FS 1"))]
        client = StubClient(error=RuntimeError("api unavailable"))
        log = logging.getLogger("progress-test-fail")

        with caplog.at_level(logging.INFO, logger="progress-test-fail"):
            describe_invoices(records, client, log)

        assert "[1/1]" in caplog.text
        assert "1 failed" in caplog.text


class TestHeaderStyling:
    """
    Header colour encodes provenance: navy for anything taken from the invoice
    or the filesystem, amber for the single column a model writes.
    """

    def sheet(self):
        return build_workbook([parse_invoice_xml(build_fa3_xml())]).active

    def test_every_header_is_bold_and_white(self):
        for cell in self.sheet()[1]:
            assert cell.font.bold is True
            assert cell.font.color.rgb.endswith(ksef_excel.HEADER_FONT_COLOUR)

    def test_invoice_sourced_headers_share_one_fill(self):
        sheet = self.sheet()
        sourced = [c for i, c in enumerate(sheet[1], start=1) if i != ksef_excel.DESCRIPTION_COLUMN]
        fills = {c.fill.fgColor.rgb for c in sourced}
        assert len(fills) == 1, "sourced headers should all use the same colour"
        assert fills.pop().endswith(ksef_excel.HEADER_FILL_SOURCED)

    def test_the_model_written_column_is_visually_distinct(self):
        sheet = self.sheet()
        generated = sheet.cell(1, ksef_excel.DESCRIPTION_COLUMN)
        sourced = sheet.cell(1, 1)

        assert generated.fill.fgColor.rgb.endswith(ksef_excel.HEADER_FILL_GENERATED)
        assert generated.fill.fgColor.rgb != sourced.fill.fgColor.rgb

    def test_headers_wrap_so_long_labels_stay_readable(self):
        for cell in self.sheet()[1]:
            assert cell.alignment.wrap_text is True

    def test_header_row_is_given_room_and_stays_frozen(self):
        sheet = self.sheet()
        assert sheet.row_dimensions[1].height == ksef_excel.HEADER_ROW_HEIGHT
        assert sheet.freeze_panes == "A2"

    def test_styling_survives_a_save_and_reload(self, tmp_path):
        target = tmp_path / "styled.xlsx"
        write_report([parse_invoice_xml(build_fa3_xml())], str(target))

        sheet = load_workbook(str(target)).active
        assert sheet.cell(1, 1).fill.fgColor.rgb.endswith(ksef_excel.HEADER_FILL_SOURCED)
        assert sheet.cell(1, ksef_excel.DESCRIPTION_COLUMN).fill.fgColor.rgb.endswith(
            ksef_excel.HEADER_FILL_GENERATED
        )
        assert sheet.cell(1, 1).font.bold is True

    def test_data_rows_are_not_given_the_header_fill(self):
        sheet = self.sheet()
        assert sheet.cell(2, 1).font.bold is not True


class TestNonDeductibleKeywords:
    """
    Keyword matching runs before any API call, so it is free, deterministic,
    and works under --no-ai.
    """

    def record(self, counterparty="Dostawca", position="Usluga hostingu"):
        from tests.fixtures.synthetic import Position

        return parse_invoice_xml(
            build_fa3_xml(seller_name=counterparty, positions=[Position(name=position)])
        )

    def test_matches_on_the_counterparty_name(self):
        assert match_non_deductible(self.record(counterparty="Benefit Systems SA"), ["benefit"])

    def test_matches_on_a_line_item(self):
        assert match_non_deductible(self.record(position="Karta Multisport Plus"), ["multisport"])

    def test_is_case_insensitive(self):
        assert match_non_deductible(self.record(position="MULTISPORT"), ["multisport"])

    def test_ignores_polish_diacritics_in_both_directions(self):
        """Invoice text arrives both ways, so matching must tolerate either."""
        assert match_non_deductible(self.record(position="odżywka białkowa"), ["odzywka"])
        assert match_non_deductible(self.record(position="odzywka bialkowa"), ["odżywka"])

    def test_returns_the_keyword_that_matched(self):
        assert (
            match_non_deductible(self.record(position="zabawka dla dziecka"), ["x", "zabawka"])
            == "zabawka"
        )

    def test_no_match_returns_none(self):
        assert match_non_deductible(self.record(), ["multisport"]) is None

    def test_empty_keyword_list_never_matches(self):
        assert match_non_deductible(self.record(position="multisport"), []) is None

    def test_marking_sets_the_label_and_clears_deductible(self, logger):
        records = [self.record(position="Karta Multisport")]

        marked = mark_non_deductible(records, ["multisport"], label="NIE PODLEGA", logger=logger)

        assert marked == 1
        assert records[0].deductible is False
        assert records[0].description == "NIE PODLEGA"

    def test_deductible_rows_are_untouched(self, logger):
        records = [self.record(position="Usluga hostingu")]

        marked = mark_non_deductible(records, ["multisport"], logger=logger)

        assert marked == 0
        assert records[0].deductible is True
        assert records[0].description == ""

    def test_the_row_is_kept_not_dropped(self, logger):
        """The whole point: nothing silently disappears from the sheet."""
        records = [self.record(position="Multisport"), self.record(position="Paliwo")]

        mark_non_deductible(records, ["multisport"], logger=logger)

        assert len(records) == 2, "a non-deductible invoice must remain in the report"


class TestCategoryClassification:
    def test_uses_the_category_returned_by_the_model(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["paliwo"])

        describe_invoices(records, client, logger, categories=["paliwo", "meble biurowe"])

        assert records[0].description == "paliwo"

    def test_the_prompt_lists_every_allowed_category(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["paliwo"])

        describe_invoices(records, client, logger, categories=["paliwo", "meble biurowe"])

        prompt = client.messages.calls[0]["messages"][0]["content"]
        assert "- paliwo" in prompt
        assert "- meble biurowe" in prompt

    def test_normalises_a_reply_back_to_the_configured_wording(self, logger):
        """A reply differing only in case or diacritics still yields the exact label."""
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["USLUGA TELEKOMUNIKACYJNA"])

        describe_invoices(records, client, logger, categories=["usługa telekomunikacyjna"])

        assert records[0].description == "usługa telekomunikacyjna"

    def test_an_off_list_reply_is_kept_and_warned_about(self, logger, caplog):
        """Discarding it would lose the classification; the warning names the gap."""
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["cos zupelnie innego"])
        log = logging.getLogger("offlist")

        with caplog.at_level(logging.WARNING, logger="offlist"):
            describe_invoices(records, client, log, categories=["paliwo"])

        assert records[0].description == "cos zupelnie innego"
        assert "not in the configured categories" in caplog.text

    def test_model_can_mark_an_invoice_non_deductible(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["NIE PODLEGA ODLICZENIU"])

        describe_invoices(records, client, logger, categories=["paliwo"])

        assert records[0].deductible is False
        assert records[0].description == "NIE PODLEGA ODLICZENIU"

    def test_already_marked_rows_cost_no_api_call(self, logger):
        """The cheap deterministic check must not be paid for twice."""
        records = [parse_invoice_xml(build_fa3_xml()) for _ in range(2)]
        records[0].deductible = False
        records[0].description = "NIE PODLEGA ODLICZENIU"
        client = StubClient(replies=["paliwo"])

        describe_invoices(records, client, logger, categories=["paliwo"])

        assert len(client.messages.calls) == 1, "the pre-marked row should be skipped"
        assert records[0].description == "NIE PODLEGA ODLICZENIU"
        assert records[1].description == "paliwo"

    def test_the_non_deductible_label_appears_in_the_prompt(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["paliwo"])

        describe_invoices(
            records, client, logger, categories=["paliwo"], non_deductible_label="POMIN"
        )

        assert "POMIN" in client.messages.calls[0]["messages"][0]["content"]

    def test_a_custom_prompt_template_is_used(self, logger):
        records = [parse_invoice_xml(build_fa3_xml())]
        client = StubClient(replies=["paliwo"])

        describe_invoices(
            records,
            client,
            logger,
            categories=["paliwo"],
            prompt_template="WLASNY {counterparty} {categories} {non_deductible_label}",
        )

        assert client.messages.calls[0]["messages"][0]["content"].startswith("WLASNY")


class TestNonDeductibleStyling:
    def test_the_marked_cell_is_highlighted_red(self):
        deductible = parse_invoice_xml(build_fa3_xml(invoice_number="FS 1"))
        deductible.description = "paliwo"
        blocked = parse_invoice_xml(build_fa3_xml(invoice_number="FS 2"))
        blocked.deductible = False
        blocked.description = "NIE PODLEGA ODLICZENIU"

        sheet = build_workbook([deductible, blocked]).active

        normal = sheet.cell(2, ksef_excel.DESCRIPTION_COLUMN)
        marked = sheet.cell(3, ksef_excel.DESCRIPTION_COLUMN)
        assert marked.font.bold is True
        assert marked.font.color.rgb.endswith(ksef_excel.NON_DEDUCTIBLE_FONT_COLOUR)
        assert marked.fill.fgColor.rgb.endswith(ksef_excel.NON_DEDUCTIBLE_FILL)
        assert normal.font.bold is not True

    def test_both_rows_are_present_in_the_sheet(self):
        blocked = parse_invoice_xml(build_fa3_xml())
        blocked.deductible = False
        sheet = build_workbook([parse_invoice_xml(build_fa3_xml()), blocked]).active
        assert sheet.max_row == 3


class TestExcelConfigDefaults:
    def test_config_without_an_excel_section_gets_defaults(self, tmp_path):
        from config_parser import load_config
        from tests.test_config_parser import MINIMAL_CONFIG, write_config

        config = load_config(write_config(tmp_path, MINIMAL_CONFIG))

        assert config.excel.categories == ksef_excel.DEFAULT_CATEGORIES
        assert "multisport" in config.excel.non_deductible_keywords
        assert config.excel.non_deductible_label == ksef_excel.DEFAULT_NON_DEDUCTIBLE_LABEL
        assert "{categories}" in config.excel.description_prompt

    def test_an_explicit_empty_keyword_list_disables_matching(self, tmp_path):
        import textwrap

        from config_parser import load_config
        from tests.test_config_parser import MINIMAL_CONFIG, write_config

        body = MINIMAL_CONFIG + textwrap.dedent("""
            excel:
              non_deductible_keywords: []
            """)
        config = load_config(write_config(tmp_path, body))
        assert config.excel.non_deductible_keywords == []

    def test_categories_and_label_can_be_overridden(self, tmp_path):
        import textwrap

        from config_parser import load_config
        from tests.test_config_parser import MINIMAL_CONFIG, write_config

        body = MINIMAL_CONFIG + textwrap.dedent("""
            excel:
              categories:
                - wlasna kategoria
              non_deductible_label: "POMIN"
            """)
        config = load_config(write_config(tmp_path, body))
        assert config.excel.categories == ["wlasna kategoria"]
        assert config.excel.non_deductible_label == "POMIN"


class TestDeduplication:
    """
    The same invoice lands on disk twice -- once as `..., ksef.pdf` from the
    download, once as a separately renamed copy. Reporting both double-counts
    the cost, so the accountant's totals come out wrong.
    """

    def record(self, number="FS 1/2026", filename="a.pdf", ksef="", date="2026-06-25"):
        r = parse_invoice_xml(build_fa3_xml(invoice_number=number, issue_date=date))
        r.filename = filename
        r.ksef_number = ksef
        return r

    def test_collapses_two_copies_of_one_invoice(self, logger):
        records = [self.record(filename="x, ksef.pdf"), self.record(filename="x, leasing.pdf")]

        result = deduplicate(records, logger=logger)

        assert len(result) == 1

    def test_different_invoices_are_both_kept(self, logger):
        records = [self.record(number="FS 1"), self.record(number="FS 2")]
        assert len(deduplicate(records, logger=logger)) == 2

    def test_same_number_different_date_is_not_merged(self, logger):
        records = [self.record(date="2026-06-25"), self.record(date="2026-07-25")]
        assert len(deduplicate(records, logger=logger)) == 2

    def test_prefers_the_copy_with_a_resolved_ksef_number(self, logger):
        number = synthetic_ksef_number()
        records = [
            self.record(filename="without.pdf", ksef=""),
            self.record(filename="with.pdf", ksef=number),
        ]

        result = deduplicate(records, logger=logger)

        assert result[0].ksef_number == number
        assert result[0].filename == "with.pdf"

    def test_prefers_a_pdf_over_an_xml_only_copy(self, logger):
        records = [self.record(filename="only.xml"), self.record(filename="real.pdf")]

        result = deduplicate(records, logger=logger)

        assert result[0].filename == "real.pdf"

    def test_the_choice_is_deterministic_regardless_of_input_order(self, logger):
        a = self.record(filename="aaa.pdf")
        b = self.record(filename="bbb.pdf")

        first = deduplicate([a, b], logger=logger)[0].filename
        second = deduplicate([b, a], logger=logger)[0].filename

        assert first == second

    def test_invoices_without_a_number_are_never_merged(self, logger):
        """An empty key would collapse unrelated invoices into one."""
        records = [
            self.record(number="", filename="a.pdf"),
            self.record(number="", filename="b.pdf"),
        ]
        assert len(deduplicate(records, logger=logger)) == 2

    def test_logs_which_file_was_dropped(self, caplog):
        log = logging.getLogger("dedup-log")
        records = [self.record(filename="kept.pdf"), self.record(filename="dropped.xml")]

        with caplog.at_level(logging.INFO, logger="dedup-log"):
            deduplicate(records, log)

        assert "kept.pdf" in caplog.text and "dropped.xml" in caplog.text
        assert "double-counted" in caplog.text

    def test_totals_are_not_double_counted(self, logger):
        """The reason this exists: the summed column must be right."""
        records = [
            self.record(number="FS 1", filename="a, ksef.pdf"),
            self.record(number="FS 1", filename="a, opis.pdf"),
            self.record(number="FS 2", filename="b, ksef.pdf"),
        ]

        result = deduplicate(records, logger=logger)

        assert sum(r.gross for r in result) == Decimal("492.00")

    def test_load_records_deduplicates_files_on_disk(self, tmp_path, logger):
        out = tmp_path / "output"
        out.mkdir(parents=True)
        xml = build_fa3_xml(invoice_number="FS 1/2026")
        for stem in ("2026-06-25, Firma, FS 1-2026, ksef", "2026-06-25, Firma, FS 1-2026, leasing"):
            (out / f"{stem}.xml").write_bytes(xml)
            (out / f"{stem}.pdf").write_bytes(b"%PDF")

        records = load_records(str(tmp_path), logger=logger)

        assert len(records) == 1, "the same invoice saved twice must yield one row"
