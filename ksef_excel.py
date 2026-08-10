#!/usr/bin/env python3
"""
Build the accountant's monthly Excel report from locally saved KSeF invoices.

The report is assembled from the FA(3) XML sidecars that ``ksef`` wrote and
``rename`` moved, never from a live KSeF query. Reading from disk keeps the
command offline, idempotent, free of API rate limits, and able to regenerate any
past month.

Each XML sidecar shares its base name with its rendered PDF, so the filename
reported in the sheet is derived from the XML path rather than matched
heuristically. When rendering failed upstream there is no PDF at all, and the
XML's own name is reported instead.

The KSeF reference number needs special handling: it is assigned by the KSeF
system and does **not** appear anywhere in the FA(3) document, so it cannot be
parsed out of the invoice. It is resolved in three steps, most authoritative
first -- the XML itself (in case a future schema carries it), the manifest that
`ksef` writes at download time, then the text of the rendered PDF, which prints
it. The last path exists to recover invoices downloaded before the manifest did.
"""

import glob
import json
import logging
import os
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from typing import Dict, List, Optional

from constants import CLAUDE_SONNET_MODEL, DEFAULT_TEMPERATURE

# A KSeF reference number looks like NIP-YYYYMMDD-XXXXXXXXXXXX-XX.
KSEF_NUMBER_PATTERN = re.compile(r"\b\d{10}-\d{8}-[0-9A-F]{12}-[0-9A-F]{2}\b", re.IGNORECASE)

# Written by `ksef` next to the invoices it saves. Keyed by invoice identity
# rather than filename, so `rename` moving and renaming files cannot break it.
MANIFEST_FILENAME = "ksef-numbers.json"

# The accountant's ten columns keep their exact headers and order so their
# existing process does not shift. "Nazwa pliku" and "pozycje" are appended.
COLUMN_HEADERS = [
    "NIP",
    "Kontrahent",
    "Numer KSeF",
    "Numer dokumentu",
    "Data wystawienia",
    "Netto",
    "VAT",
    "Brutto",
    "Waluta",
    "Opis faktury - czego dotyczy w działalnosci",
    "Nazwa pliku",
    "pozycje",
]

DESCRIPTION_COLUMN = 10
POZYCJE_COLUMN = 12
MONEY_COLUMNS = (6, 7, 8)
MONEY_FORMAT = "#,##0.00"
MAX_TOKENS_DESCRIPTION = 200

# Header colours encode provenance rather than decorate: everything taken from
# the invoice or the filesystem is navy, and the one column written by a model
# is amber, so the accountant can see at a glance which cells to sanity-check.
HEADER_FILL_SOURCED = "1F3864"
HEADER_FILL_GENERATED = "7F6000"
HEADER_FONT_COLOUR = "FFFFFF"
HEADER_ROW_HEIGHT = 32

DESCRIPTION_PROMPT = """Jesteś księgowym. Na podstawie danych faktury zakupowej napisz
krótkie uzasadnienie, czego dotyczy wydatek w działalności gospodarczej — tak, aby
posłużyło jako podstawa odliczenia podatkowego.

Kontrahent: {counterparty}
Kwoty: netto {net}, VAT {vat}, brutto {gross} {currency}
Pozycje na fakturze:
{positions}

Odpowiedz jednym zdaniem po polsku, bez wstępu i bez cudzysłowów. Maksymalnie 200 znaków."""


@dataclass
class InvoicePosition:
    """One line item from an FA(3) ``FaWiersz`` element."""

    name: str
    unit: str
    quantity: Decimal
    unit_net: Decimal
    net: Decimal
    vat_rate: str

    @property
    def vat(self) -> Decimal:
        """VAT for this line, or zero when the rate is a code such as 'zw'."""
        if not _is_numeric_rate(self.vat_rate):
            return Decimal("0.00")
        value = self.net * Decimal(self.vat_rate) / Decimal(100)
        return value.quantize(Decimal("0.01"))

    @property
    def gross(self) -> Decimal:
        return self.net + self.vat


@dataclass
class InvoiceRecord:
    """One row of the report."""

    nip: str = ""
    counterparty: str = ""
    # Both parties are kept regardless of role: the manifest is keyed on the
    # seller, while the report shows whichever party is the counterparty.
    seller_nip: str = ""
    buyer_nip: str = ""
    ksef_number: str = ""
    document_number: str = ""
    issue_date: str = ""
    net: Decimal = Decimal("0.00")
    vat: Decimal = Decimal("0.00")
    gross: Decimal = Decimal("0.00")
    currency: str = ""
    description: str = ""
    filename: str = ""
    positions: List[InvoicePosition] = field(default_factory=list)


def _local_name(tag: str) -> str:
    """Strip any namespace from an element tag."""
    return tag.rsplit("}", 1)[-1] if "}" in tag else tag


def _first(element: ET.Element, name: str) -> Optional[ET.Element]:
    """Find the first descendant (or self) whose local name matches."""
    if _local_name(element.tag) == name:
        return element
    for child in element.iter():
        if _local_name(child.tag) == name:
            return child
    return None


def _children_matching(element: Optional[ET.Element], pattern: str) -> List[ET.Element]:
    """All descendants whose local name matches a regular expression."""
    if element is None:
        return []
    regex = re.compile(pattern)
    return [child for child in element.iter() if regex.fullmatch(_local_name(child.tag))]


def _text(element: Optional[ET.Element], name: str) -> str:
    """Text of the first matching descendant, stripped, or empty string."""
    if element is None:
        return ""
    found = _first(element, name)
    if found is None or found.text is None:
        return ""
    return found.text.strip()


def _is_numeric_rate(rate: str) -> bool:
    try:
        Decimal(rate)
        return True
    except (InvalidOperation, ValueError, ArithmeticError):
        return False


def _to_decimal(text: str, default: str = "0.00") -> Decimal:
    """
    Parse an FA(3) monetary value.

    KSeF uses a dot decimal separator, but this tolerates a comma and embedded
    spaces so an oddly serialised document does not lose its amounts.
    """
    if not text:
        return Decimal(default)
    cleaned = text.replace(" ", "").replace(" ", "").replace(",", ".")
    try:
        return Decimal(cleaned)
    except (InvalidOperation, ValueError, ArithmeticError):
        return Decimal(default)


def _party(root: ET.Element, element_name: str) -> tuple:
    """Return (name, nip) for a Podmiot1/Podmiot2 element."""
    party = _first(root, element_name)
    if party is None:
        return "", ""
    return _text(party, "Nazwa"), _text(party, "NIP")


def parse_invoice_xml(xml_bytes: bytes, role: str = "buyer") -> InvoiceRecord:
    """
    Parse one FA(3) document into a report row.

    Args:
        xml_bytes: Raw FA(3) XML.
        role: ``"buyer"`` reports the seller as counterparty (purchase
            invoices); ``"seller"`` reports the buyer.

    Raises:
        ValueError: If the document is not parseable XML.
    """
    try:
        root = ET.fromstring(xml_bytes)
    except ET.ParseError as e:
        raise ValueError(f"Not a parseable XML document: {e}") from e

    fa = _first(root, "Fa")

    seller_name, seller_nip = _party(root, "Podmiot1")
    buyer_name, buyer_nip = _party(root, "Podmiot2")
    counterparty, nip = (seller_name, seller_nip) if role == "buyer" else (buyer_name, buyer_nip)

    positions = [
        InvoicePosition(
            name=_text(row, "P_7"),
            unit=_text(row, "P_8A"),
            quantity=_to_decimal(_text(row, "P_8B")),
            unit_net=_to_decimal(_text(row, "P_9A")),
            net=_to_decimal(_text(row, "P_11")),
            vat_rate=_text(row, "P_12"),
        )
        for row in _children_matching(fa, r"FaWiersz")
    ]

    # Net and VAT totals are reported per VAT rate in numbered field pairs, so
    # every P_13_x / P_14_x must be summed -- not just the first.
    net_elements = _children_matching(fa, r"P_13_\d+")
    vat_elements = _children_matching(fa, r"P_14_\d+")

    if net_elements:
        net = sum((_to_decimal(e.text or "") for e in net_elements), Decimal("0.00"))
    else:
        # No per-rate net totals: recover it from the line items, which carry
        # their own net values. Leaving net at zero would make the VAT fallback
        # below attribute the entire gross amount to tax.
        net = sum((p.net for p in positions), Decimal("0.00"))

    vat = sum((_to_decimal(e.text or "") for e in vat_elements), Decimal("0.00"))

    gross_text = _text(fa, "P_15")
    gross = _to_decimal(gross_text) if gross_text else net + vat

    # Some documents omit the per-rate pairs entirely; recover VAT from the
    # difference rather than reporting zero.
    if not vat_elements:
        vat = gross - net

    return InvoiceRecord(
        nip=nip,
        counterparty=counterparty,
        seller_nip=seller_nip,
        buyer_nip=buyer_nip,
        # Real FA(3) documents do not carry the KSeF reference number -- it is
        # assigned by the system and returned in metadata. This reads it when
        # present, and load_records recovers it from the manifest or the
        # rendered PDF when it is not.
        ksef_number=_text(root, "NumerKSeFDokumentu"),
        document_number=_text(fa, "P_2"),
        issue_date=_text(fa, "P_1"),
        net=net,
        vat=vat,
        gross=gross,
        currency=_text(fa, "KodWaluty"),
        description="",
        filename="",
        positions=positions,
    )


def _format_quantity(quantity: Decimal) -> str:
    """Render a quantity without trailing zeros: 1.000 becomes 1."""
    normalized = quantity.normalize()
    if normalized == normalized.to_integral_value():
        try:
            return str(normalized.quantize(Decimal("1")))
        except InvalidOperation:
            return str(normalized)
    return str(normalized)


def _format_money(value: Decimal) -> str:
    return f"{value:.2f}"


def format_positions(positions: List[InvoicePosition]) -> str:
    """
    Flatten line items into a single cell, one numbered entry per line.

    Twelve columns wide, a single run-on line would be unreadable, so entries
    are newline-separated and the cell is set to wrap.
    """
    lines = []
    for number, position in enumerate(positions, start=1):
        rate = f"{position.vat_rate}%" if _is_numeric_rate(position.vat_rate) else position.vat_rate
        lines.append(
            f"{number}. {position.name} — "
            f"{_format_quantity(position.quantity)} {position.unit} "
            f"× {_format_money(position.unit_net)} = "
            f"netto {_format_money(position.net)}, "
            f"VAT {rate}, "
            f"brutto {_format_money(position.gross)}"
        )
    return "\n".join(lines)


def format_date(value: str) -> str:
    """Convert an ISO date to the DD.MM.YYYY order the accountant's sheet uses."""
    if not value:
        return ""
    match = re.match(r"^(\d{4})-(\d{2})-(\d{2})", value)
    if not match:
        return value
    year, month, day = match.groups()
    return f"{day}.{month}.{year}"


def manifest_key(seller_nip: str, invoice_number: str, issue_date: str) -> str:
    """
    Build the manifest key identifying an invoice by content, not filename.

    Keying on identity rather than path is what lets `rename` move and rename
    files without orphaning their KSeF numbers.
    """
    return f"{seller_nip.strip()}|{invoice_number.strip()}|{issue_date.strip()}"


def load_ksef_manifest(directory: str, logger: Optional[logging.Logger] = None) -> Dict[str, str]:
    """
    Merge every ``ksef-numbers.json`` found under ``directory``.

    `ksef` writes one manifest per month directory, so a report spanning
    several directories merges several manifests.
    """
    log = logger or logging.getLogger(__name__)
    merged: Dict[str, str] = {}

    pattern = os.path.join(directory, "**", MANIFEST_FILENAME)
    for path in sorted(glob.glob(pattern, recursive=True)):
        try:
            with open(path, "r", encoding="utf-8") as handle:
                entries = json.load(handle)
            if isinstance(entries, dict):
                merged.update(entries)
                log.debug(f"Read {len(entries)} KSeF number(s) from {path}")
        except (OSError, json.JSONDecodeError) as e:
            log.warning(f"Could not read KSeF manifest '{path}': {e}")

    return merged


def extract_ksef_number_from_pdf(pdf_path: str) -> Optional[str]:
    """
    Recover the KSeF reference number from a rendered PDF's text.

    The renderer prints the number on the page, so this recovers it for
    invoices downloaded before the manifest existed. Returns None if the file
    is unreadable or contains no number-shaped string.
    """
    try:
        import PyPDF2

        reader = PyPDF2.PdfReader(pdf_path)
        text = "".join((page.extract_text() or "") for page in reader.pages[:2])
    except Exception:
        # PyPDF2 raises a wide range of errors on damaged files; none of them
        # should stop the report.
        return None

    match = KSEF_NUMBER_PATTERN.search(text)
    return match.group(0) if match else None


def _resolve_filename(xml_path: str, log: logging.Logger) -> str:
    """
    Report the file the accountant can actually open.

    Prefers the sibling PDF. When rendering failed upstream there is no PDF at
    all, so naming a non-existent file would be worse than naming the XML that
    does exist.
    """
    pdf_path = os.path.splitext(xml_path)[0] + ".pdf"
    if os.path.exists(pdf_path):
        return os.path.basename(pdf_path)
    log.debug(f"No PDF beside '{os.path.basename(xml_path)}'; reporting the XML name.")
    return os.path.basename(xml_path)


def load_records(
    directory: str,
    role: str = "buyer",
    logger: Optional[logging.Logger] = None,
) -> List[InvoiceRecord]:
    """
    Load every FA(3) sidecar under ``directory``, recursively.

    Searching recursively rather than assuming a fixed layout means the command
    works whether invoices sit in ``invoices/output/`` or a month subdirectory.

    A document that fails to parse is skipped with a warning: one bad invoice
    must not cost the accountant the other twenty rows.

    The KSeF number is resolved in three steps, most authoritative first: the
    XML itself, then the manifest written at download time, then the text of
    the rendered PDF.
    """
    log = logger or logging.getLogger(__name__)

    if not os.path.isdir(directory):
        log.warning(f"Directory not found, nothing to report: {directory}")
        return []

    manifest = load_ksef_manifest(directory, logger=log)
    if manifest:
        log.info(f"Loaded {len(manifest)} KSeF number(s) from manifest files.")

    xml_paths = sorted(glob.glob(os.path.join(directory, "**", "*.xml"), recursive=True))
    log.info(f"Scanning {len(xml_paths)} XML file(s) under {directory}")

    records: List[InvoiceRecord] = []
    sources = {"xml": 0, "manifest": 0, "pdf": 0, "missing": 0}
    skipped = 0
    without_pdf = 0

    for index, xml_path in enumerate(xml_paths, start=1):
        name = os.path.basename(xml_path)
        try:
            with open(xml_path, "rb") as handle:
                record = parse_invoice_xml(handle.read(), role=role)
        except (ValueError, OSError) as e:
            skipped += 1
            log.warning(f"[{index}/{len(xml_paths)}] Skipping '{name}': {e}")
            continue

        pdf_path = os.path.splitext(xml_path)[0] + ".pdf"
        if not os.path.exists(pdf_path):
            without_pdf += 1
        record.filename = _resolve_filename(xml_path, log)

        if record.ksef_number:
            source = "xml"
        else:
            key = manifest_key(record.seller_nip, record.document_number, record.issue_date)
            from_manifest = manifest.get(key)
            if from_manifest:
                record.ksef_number = from_manifest
                source = "manifest"
            else:
                from_pdf = extract_ksef_number_from_pdf(pdf_path)
                if from_pdf:
                    record.ksef_number = from_pdf
                    source = "pdf"
                else:
                    source = "missing"
        sources[source] += 1

        log.debug(
            f"[{index}/{len(xml_paths)}] {record.issue_date} "
            f"{record.document_number or '(no number)'} "
            f"-- KSeF number from {source}"
        )
        records.append(record)

    records.sort(key=lambda r: (r.issue_date, r.document_number))

    log.info(
        f"Loaded {len(records)} invoice(s); KSeF number from XML {sources['xml']}, "
        f"manifest {sources['manifest']}, PDF {sources['pdf']}, "
        f"unresolved {sources['missing']}"
    )
    if skipped:
        log.warning(f"{skipped} file(s) skipped as unparseable (run with --log-level DEBUG).")
    if without_pdf:
        log.warning(
            f"{without_pdf} invoice(s) have no rendered PDF; the report names their "
            f"XML instead (PDF rendering failed when they were downloaded)."
        )
    if sources["missing"]:
        log.warning(
            f"{sources['missing']} invoice(s) have no KSeF number available from any "
            f"source; re-run `ksef` for those months to record it."
        )
    return records


def filter_by_month(records: List[InvoiceRecord], month: Optional[str]) -> List[InvoiceRecord]:
    """
    Keep only records issued in ``month`` (``YYYY-MM``). ``None`` keeps all.

    Filtering on the issue date in the document, rather than on directory
    layout, keeps the result correct however the files happen to be arranged.
    """
    if not month:
        return list(records)
    return [r for r in records if r.issue_date.startswith(month)]


def build_workbook(records: List[InvoiceRecord]):
    """Assemble the report workbook. Requires no API access."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Faktury"

    sheet.append(COLUMN_HEADERS)

    sourced_fill = PatternFill("solid", fgColor=HEADER_FILL_SOURCED)
    generated_fill = PatternFill("solid", fgColor=HEADER_FILL_GENERATED)
    header_font = Font(bold=True, color=HEADER_FONT_COLOUR, size=11)
    header_border = Border(bottom=Side(style="thin", color=HEADER_FILL_SOURCED))

    for column, cell in enumerate(sheet[1], start=1):
        cell.font = header_font
        cell.fill = generated_fill if column == DESCRIPTION_COLUMN else sourced_fill
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        cell.border = header_border
    sheet.row_dimensions[1].height = HEADER_ROW_HEIGHT

    for record in records:
        sheet.append(
            [
                record.nip,
                record.counterparty,
                record.ksef_number,
                record.document_number,
                format_date(record.issue_date),
                float(record.net),
                float(record.vat),
                float(record.gross),
                record.currency,
                record.description,
                record.filename,
                format_positions(record.positions),
            ]
        )
        row = sheet.max_row
        for column in MONEY_COLUMNS:
            sheet.cell(row, column).number_format = MONEY_FORMAT
        sheet.cell(row, POZYCJE_COLUMN).alignment = Alignment(wrap_text=True, vertical="top")
        sheet.cell(row, DESCRIPTION_COLUMN).alignment = Alignment(wrap_text=True, vertical="top")

    widths = [14, 34, 40, 20, 16, 12, 12, 12, 8, 46, 44, 60]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width

    sheet.freeze_panes = "A2"
    return workbook


def write_report(records: List[InvoiceRecord], path: str) -> None:
    """Write the workbook to ``path``, creating parent directories as needed."""
    parent = os.path.dirname(os.path.abspath(path))
    os.makedirs(parent, exist_ok=True)
    build_workbook(records).save(path)


def describe_invoices(
    records: List[InvoiceRecord],
    client,
    logger: logging.Logger,
    model: str = CLAUDE_SONNET_MODEL,
) -> None:
    """
    Fill each record's ``description`` with a tax-deduction rationale.

    Sends only the counterparty, line items, and totals -- deliberately not the
    full XML, which carries more than the task requires.

    Failures are per-row: a failed call leaves that description empty, logs a
    warning, and never aborts the remaining invoices.

    Progress is logged per invoice. This runs one sequential API call per row,
    so on a busy month it is the slowest step by far -- without a line per item
    it is indistinguishable from a hang.
    """
    total = len(records)
    failures = 0

    for index, record in enumerate(records, start=1):
        label = f"[{index}/{total}]"
        logger.info(
            f"{label} Describing {record.document_number or '(no number)'} "
            f"from {record.counterparty or 'unknown'}"
        )
        prompt = DESCRIPTION_PROMPT.format(
            counterparty=record.counterparty or "nieznany",
            net=_format_money(record.net),
            vat=_format_money(record.vat),
            gross=_format_money(record.gross),
            currency=record.currency or "PLN",
            positions=format_positions(record.positions) or "(brak pozycji)",
        )
        try:
            message = client.messages.create(
                model=model,
                max_tokens=MAX_TOKENS_DESCRIPTION,
                temperature=DEFAULT_TEMPERATURE,
                messages=[{"role": "user", "content": prompt}],
            )
            record.description = message.content[0].text.strip()
            logger.info(f"{label} -> {record.description}")
        except Exception as e:
            # Deliberately broad: an SDK, network, or content error on one
            # invoice must not cost the rest of the report.
            failures += 1
            logger.warning(f"{label} Description failed for '{record.document_number}': {e}")
            record.description = ""

    if total:
        logger.info(f"Descriptions complete: {total - failures} written, {failures} failed.")
